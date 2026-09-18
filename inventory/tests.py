import io
import json
import shutil
import tempfile
from datetime import timedelta

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from django.utils import timezone
from openpyxl import load_workbook
from PIL import Image
from rest_framework import status
from rest_framework.test import APITestCase

from .models import CODE_ALPHABET, CODE_LENGTH, Batch, Restock, Sale, SizeEntry

MEDIA = tempfile.mkdtemp()


@override_settings(MEDIA_ROOT=MEDIA)
class ApiTestCase(APITestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(MEDIA, ignore_errors=True)

    def setUp(self):
        self.user = get_user_model().objects.create_user('owner', password='bazaar-pass-2026')
        self.client.force_authenticate(self.user)

    def add(self, brand='Nike Air', price=250_000, sizes=(('41', 5), ('42', 3))):
        body = {'brand': brand, 'bought_price': price, 'sizes': [{'size': s, 'quantity': q} for s, q in sizes]}
        return self.client.post('/api/batches/', body, format='json')

    def sell(self, code, price=300_000):
        return self.client.post(f'/api/entries/{code}/sell/', {'sold_price': price}, format='json')

    def entry(self, size='41', brand='Nike Air'):
        return SizeEntry.objects.get(size=size, batch__brand=brand)


class AuthTests(ApiTestCase):
    def test_endpoints_require_a_token(self):
        self.client.force_authenticate(None)
        for url in ['/api/entries/', '/api/batches/', '/api/sales/', '/api/stats/', '/api/export/inventory.csv']:
            self.assertEqual(self.client.get(url).status_code, status.HTTP_401_UNAUTHORIZED, url)

    def test_login_returns_jwt_pair_that_works(self):
        self.client.force_authenticate(None)
        response = self.client.post('/api/auth/token/', {'username': 'owner', 'password': 'bazaar-pass-2026'}, format='json')
        self.assertEqual(response.status_code, 200)
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {response.data['access']}")
        self.assertEqual(self.client.get('/api/auth/me/').data['username'], 'owner')

        refreshed = self.client.post('/api/auth/token/refresh/', {'refresh': response.data['refresh']}, format='json')
        self.assertEqual(refreshed.status_code, 200)
        self.assertIn('access', refreshed.data)

    def test_wrong_password_is_rejected(self):
        self.client.force_authenticate(None)
        response = self.client.post('/api/auth/token/', {'username': 'owner', 'password': 'nope'}, format='json')
        self.assertEqual(response.status_code, 401)


class StockEntryTests(ApiTestCase):
    def test_new_batch_creates_one_line_per_size_with_unique_codes(self):
        response = self.add()
        self.assertEqual(response.status_code, 201)
        self.assertEqual(Batch.objects.count(), 1)
        lines = response.data['lines']
        self.assertEqual([(l['entry']['size'], l['added'], l['merged']) for l in lines], [('41', 5, False), ('42', 3, False)])
        codes = {l['entry']['code'] for l in lines}
        self.assertEqual(len(codes), 2)
        for code in codes:
            self.assertEqual(len(code), CODE_LENGTH)
            self.assertLessEqual(set(code), set(CODE_ALPHABET))

    def test_same_brand_size_and_price_tops_up_existing_line(self):
        self.add()
        code = self.entry('41').code
        response = self.add(brand='  nike   air ', sizes=(('41', 4), ('43', 2)))
        self.assertEqual(response.status_code, 201)
        merged = next(l for l in response.data['lines'] if l['merged'])
        self.assertEqual((merged['entry']['code'], merged['added'], merged['entry']['quantity']), (code, 4, 9))
        entry = self.entry('41')
        self.assertEqual((entry.quantity, entry.initial_quantity), (9, 9))
        self.assertEqual(Batch.objects.count(), 2)  # size 43 was new, so it needed a batch
        self.assertEqual(SizeEntry.objects.count(), 3)
        self.assertEqual(set(Batch.objects.values_list('brand', flat=True)), {'Nike Air'})

    def test_all_sizes_merged_creates_no_empty_batch(self):
        self.add()
        response = self.add(sizes=(('41', 1),))
        self.assertIsNone(response.data['batch'])
        self.assertEqual(Batch.objects.count(), 1)

    def test_different_price_gets_a_new_code(self):
        self.add()
        self.add(price=270_000, sizes=(('41', 2),))
        self.assertEqual(SizeEntry.objects.filter(size='41').count(), 2)

    def test_sold_out_line_is_not_topped_up(self):
        self.add(sizes=(('41', 1),))
        self.sell(self.entry('41').code)
        self.add(sizes=(('41', 2),))
        self.assertEqual(SizeEntry.objects.filter(size='41').count(), 2)

    def test_repeated_sizes_in_one_request_are_summed(self):
        self.add(sizes=(('41', 2), (' 41 ', 3), ('41,5', 1)))
        self.assertEqual(self.entry('41').quantity, 5)
        self.assertEqual(SizeEntry.objects.get(size='41.5').quantity, 1)

    def test_validation_errors_create_nothing(self):
        for response in [
            self.add(sizes=()),
            self.add(price=0),
            self.add(sizes=(('XL', 1),)),
            self.add(sizes=(('41', 2), ('41.0', 1))),
            self.add(sizes=(('41', 0),)),
            self.add(brand='   '),
        ]:
            self.assertEqual(response.status_code, 400, response.data)
        self.assertEqual(SizeEntry.objects.count(), 0)

    def test_multipart_with_photo_and_sizes_as_json_string(self):
        image = io.BytesIO()
        Image.new('RGB', (20, 20), 'navy').save(image, format='JPEG')
        response = self.client.post('/api/batches/', {
            'brand': 'Adidas', 'bought_price': '180000',
            'sizes': json.dumps([{'size': '40', 'quantity': 2}]),
            'picture': SimpleUploadedFile('IMG_0001.JPG', image.getvalue(), content_type='image/jpeg'),
        }, format='multipart')
        self.assertEqual(response.status_code, 201, response.data)
        picture = Batch.objects.get().picture.name
        self.assertTrue(picture.startswith('batches/') and picture.endswith('.jpg'))
        self.assertNotIn('IMG_0001', picture)

    def test_brands_autocomplete(self):
        self.add(brand='Puma')
        self.add(brand='Adidas')
        self.assertEqual(self.client.get('/api/batches/brands/').data, ['Adidas', 'Puma'])


class EntryTests(ApiTestCase):
    def setUp(self):
        super().setUp()
        self.add()
        self.add(brand='Adidas', price=180_000, sizes=(('40', 1),))

    def codes(self, query):
        return [e['code'] for e in self.client.get(f'/api/entries/?{query}').data['results']]

    def test_detail_by_code_is_case_insensitive_and_includes_barcode(self):
        code = self.entry('41').code
        response = self.client.get(f'/api/entries/{code.lower()}/')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['code'], code)
        self.assertTrue(response.data['barcode_svg'].startswith('<svg'))
        self.assertEqual(response.data['batch']['bought_price'], 250_000)

    def test_unknown_code_is_404(self):
        self.assertEqual(self.client.get('/api/entries/ZZZZZZZZZZ/').status_code, 404)

    def test_filters_and_ordering(self):
        adidas = self.entry('40', 'Adidas').code
        self.assertEqual(len(self.codes('q=nike')), 2)
        self.assertEqual(self.codes(f'q={adidas.lower()}'), [adidas])
        self.assertEqual(len(self.codes('size=42')), 1)
        self.sell(adidas, 1)
        self.assertEqual(len(self.codes('stock=in')), 2)
        self.assertEqual(self.codes('stock=out'), [adidas])
        quantities = [e['quantity'] for e in self.client.get('/api/entries/?ordering=-quantity').data['results']]
        self.assertEqual(quantities, sorted(quantities, reverse=True))

    def test_correcting_quantity_does_not_count_as_sold(self):
        code = self.entry('41').code
        response = self.client.patch(f'/api/entries/{code}/', {'quantity': 3}, format='json')
        self.assertEqual(response.status_code, 200)
        self.assertEqual((response.data['quantity'], response.data['initial_quantity'], response.data['sold']), (3, 3, 0))

    def test_code_and_size_are_read_only(self):
        entry = self.entry('41')
        self.client.patch(f'/api/entries/{entry.code}/', {'code': 'AAAAAAAAAA', 'size': '45'}, format='json')
        entry.refresh_from_db()
        self.assertEqual(entry.size, '41')
        self.assertNotEqual(entry.code, 'AAAAAAAAAA')

    def test_cannot_delete_line_with_sales(self):
        code = self.entry('41').code
        self.sell(code)
        self.assertEqual(self.client.delete(f'/api/entries/{code}/').status_code, 409)
        self.assertEqual(self.client.delete(f"/api/entries/{self.entry('42').code}/").status_code, 204)


class SellTests(ApiTestCase):
    def setUp(self):
        super().setUp()
        self.add(sizes=(('41', 2),))
        self.code = self.entry('41').code

    def test_sell_records_sale_and_decrements(self):
        response = self.sell(self.code)
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.data['entry']['quantity'], 1)
        self.assertEqual(response.data['sale']['profit'], 50_000)
        self.assertEqual(Sale.objects.count(), 1)

    def test_selling_below_cost_is_allowed_and_shows_loss(self):
        self.assertEqual(self.sell(self.code, 200_000).data['sale']['profit'], -50_000)

    def test_cannot_sell_past_zero(self):
        self.sell(self.code)
        self.sell(self.code)
        self.assertEqual(self.sell(self.code).status_code, 409)
        self.assertEqual(self.entry('41').quantity, 0)
        self.assertEqual(Sale.objects.count(), 2)

    def test_invalid_price(self):
        self.assertEqual(self.sell(self.code, -1).status_code, 400)
        self.assertEqual(self.client.post(f'/api/entries/{self.code}/sell/', {}, format='json').status_code, 400)
        self.assertEqual(Sale.objects.count(), 0)

    def test_sales_cannot_be_deleted_through_the_api(self):
        sale_id = self.sell(self.code).data['sale']['id']
        self.assertEqual(self.client.delete(f'/api/sales/{sale_id}/').status_code, 405)
        self.assertEqual(Sale.objects.count(), 1)

    def test_admin_delete_restocks_the_pair(self):
        sale_id = self.sell(self.code).data['sale']['id']
        admin = get_user_model().objects.create_superuser('boss', password='admin-pass-2026')
        self.client.force_login(admin)
        response = self.client.post(f'/admin/inventory/sale/{sale_id}/delete/', {'post': 'yes'})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(self.entry('41').quantity, 2)
        self.assertEqual(Sale.objects.count(), 0)

    def test_last_brand_price_suggestion(self):
        self.sell(self.code, 310_000)
        self.assertEqual(self.client.get(f'/api/entries/{self.code}/').data['last_brand_price'], 310_000)


class LabelTests(ApiTestCase):
    def setUp(self):
        super().setUp()
        self.add(brand='Кроссовки Ecco', sizes=(('41', 2), ('42', 1)))

    def test_single_label_png(self):
        entry = SizeEntry.objects.get(size='41')
        response = self.client.get(f'/api/entries/{entry.code}/label/')
        self.assertEqual(response['Content-Type'], 'image/png')
        self.assertEqual(Image.open(io.BytesIO(response.content)).size, (709, 354))

    def test_pdf_sheet_spans_pages_and_marks_lines_printed(self):
        codes = list(SizeEntry.objects.values_list('code', flat=True))
        response = self.client.post('/api/labels/pdf/', {'items': [{'code': c, 'copies': 30} for c in codes]}, format='json')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertTrue(response.content.startswith(b'%PDF'))
        self.assertIn(b'/Count 3', response.content)  # 60 labels, 3x9 = 27 per A4 page
        self.assertFalse(SizeEntry.objects.filter(label_printed=False).exists())

    def test_pdf_rejects_unknown_codes(self):
        response = self.client.post('/api/labels/pdf/', {'items': [{'code': 'ZZZZZZZZZZ'}]}, format='json')
        self.assertEqual(response.status_code, 400)


class StatsTests(ApiTestCase):
    def test_numbers(self):
        self.add(sizes=(('41', 3), ('42', 2)))                        # 5 pairs x 250k
        self.add(brand='Adidas', price=100_000, sizes=(('40', 4),))   # 4 pairs x 100k
        nike, adidas = self.entry('41').code, self.entry('40', 'Adidas').code
        self.sell(nike, 300_000)
        self.sell(nike, 280_000)
        self.sell(adidas, 150_000)

        data = self.client.get('/api/stats/').data
        self.assertEqual(data['inventory'], {'value': 3 * 250_000 + 3 * 100_000, 'pairs': 6})
        self.assertEqual(data['sales'], {'revenue': 730_000, 'cost': 600_000, 'profit': 130_000, 'units': 3})
        self.assertEqual((data['best_brands'][0]['label'], data['best_brands'][0]['units']), ('Nike Air', 2))
        self.assertEqual(data['best_sizes'][0]['label'], '41')
        self.assertEqual(data['timeline']['step'], 'day')
        self.assertEqual(len(data['timeline']['points']), 30)
        self.assertEqual(data['timeline']['points'][-1]['revenue'], 730_000)
        self.assertEqual(data['average_days_to_sell'], 0.0)

    def test_slow_moving_stock(self):
        self.add(sizes=(('41', 3), ('42', 2)))
        Batch.objects.update(date_added=timezone.now() - timedelta(days=60))
        self.sell(self.entry('42').code, 1)
        slow = self.client.get('/api/stats/?slow_days=30').data['slow_moving']
        self.assertEqual([e['size'] for e in slow['entries']], ['41'])

    def test_bad_range(self):
        self.assertEqual(self.client.get('/api/stats/?start=2026-09-10&end=2026-09-01').status_code, 400)

    def test_long_range_groups_by_month(self):
        self.assertEqual(self.client.get('/api/stats/?start=2025-01-01&end=2025-12-31').data['timeline']['step'], 'month')


class ExportTests(ApiTestCase):
    def setUp(self):
        super().setUp()
        self.add(brand='Обувь', sizes=(('41', 2),))
        self.sell(self.entry('41', 'Обувь').code)

    def test_inventory_csv_has_bom_and_cyrillic(self):
        response = self.client.get('/api/export/inventory.csv')
        self.assertEqual(response.status_code, 200)
        text = response.content.decode('utf-8')
        self.assertTrue(text.startswith('﻿'))
        self.assertIn('Обувь', text)
        self.assertIn('attachment; filename="inventory-', response['Content-Disposition'])

    def test_sales_xlsx(self):
        response = self.client.get('/api/export/sales.xlsx')
        sheet = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet.cell(2, 7).value, 50_000)

    def test_unknown_export_is_404(self):
        self.assertEqual(self.client.get('/api/export/users.csv').status_code, 404)


class DailyReportTests(ApiTestCase):
    def setUp(self):
        super().setUp()
        self.today = timezone.localdate()
        self.yesterday = self.today - timedelta(days=1)
        self.add(sizes=(('41', 5),))                                   # 5 x 250k received today
        self.add(brand='Adidas', price=100_000, sizes=(('40', 2),))    # 2 x 100k received yesterday
        Restock.objects.filter(size_entry__batch__brand='Adidas').update(received_at=timezone.now() - timedelta(days=1))
        self.nike, self.adidas = self.entry('41').code, self.entry('40', 'Adidas').code
        self.sell(self.nike, 300_000)
        self.sell(self.nike, 240_000)
        old = self.sell(self.adidas, 150_000).data['sale']['id']
        Sale.objects.filter(pk=old).update(sold_at=timezone.now() - timedelta(days=1))

    def test_today(self):
        data = self.client.get('/api/reports/daily/').data
        self.assertEqual(data['date'], self.today)
        self.assertEqual(data['sales'], {'revenue': 540_000, 'cost': 500_000, 'profit': 40_000, 'units': 2})
        self.assertEqual([s['profit'] for s in data['sales_list']], [50_000, -10_000])
        self.assertEqual(data['by_brand'][0]['units'], 2)
        self.assertEqual(data['received_summary'], {'pairs': 5, 'value': 1_250_000})

    def test_top_up_counts_as_received_on_its_own_day(self):
        self.add(sizes=(('41', 3),))  # merges into the existing line
        data = self.client.get('/api/reports/daily/').data
        self.assertEqual(data['received_summary']['pairs'], 8)
        self.assertEqual(len(data['received']), 2)

    def test_specific_day(self):
        data = self.client.get(f'/api/reports/daily/?date={self.yesterday}').data
        self.assertEqual(data['sales']['units'], 1)
        self.assertEqual(data['received_summary'], {'pairs': 2, 'value': 200_000})

    def test_days_list_includes_quiet_days(self):
        data = self.client.get(f'/api/reports/days/?start={self.today - timedelta(days=6)}&end={self.today}').data
        self.assertEqual(len(data['days']), 7)
        self.assertEqual(data['days'][0], {
            'date': self.today, 'units': 2, 'revenue': 540_000, 'profit': 40_000,
            'received_pairs': 5, 'received_value': 1_250_000,
        })
        self.assertEqual((data['days'][1]['units'], data['days'][1]['received_pairs']), (1, 2))
        self.assertEqual(data['days'][2]['units'], 0)

    def test_days_range_limit_and_bad_date(self):
        self.assertEqual(self.client.get('/api/reports/days/?start=2020-01-01&end=2026-01-01').status_code, 400)
        self.assertEqual(self.client.get('/api/reports/daily/?date=yesterday').status_code, 400)

    def test_excel_export(self):
        response = self.client.get(f'/api/reports/daily.xlsx?date={self.today}')
        self.assertEqual(response.status_code, 200)
        self.assertIn(f'daily-report-{self.today}', response['Content-Disposition'])
        book = load_workbook(io.BytesIO(response.content))
        sales, received = book.worksheets
        self.assertEqual(sales.max_row, 4)  # header, 2 sales, total
        self.assertEqual(sales.cell(4, 7).value, 40_000)
        self.assertEqual(received.cell(received.max_row, 6).value, 1_250_000)

    def test_backfilled_restocks_exist_for_new_stock(self):
        self.assertEqual(Restock.objects.count(), 2)
