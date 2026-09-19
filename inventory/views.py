import csv
import io
from collections import defaultdict
from datetime import timedelta

from django.db import transaction
from django.db.models import Count, F, Max, Min, ProtectedError, Sum
from django.http import HttpResponse
from django.utils import timezone
from django.utils.translation import gettext as _
from drf_spectacular.types import OpenApiTypes
from drf_spectacular.utils import OpenApiParameter, OpenApiResponse, extend_schema
from rest_framework import mixins, status, viewsets
from rest_framework.decorators import action
from rest_framework.exceptions import ValidationError
from rest_framework.parsers import FormParser, JSONParser, MultiPartParser
from rest_framework.response import Response
from rest_framework.throttling import ScopedRateThrottle
from rest_framework.views import APIView
from rest_framework_simplejwt.views import TokenObtainPairView

from . import reports, services, stats
from .filters import SaleFilter, SizeEntryFilter
from .models import Batch, Sale, SizeEntry
from .serializers import (
    BatchCreateSerializer, BatchSerializer, BrandGroupSerializer, DateRangeSerializer, DayQuerySerializer, LabelSheetSerializer,
    RestockResultSerializer, SaleSerializer, SellResultSerializer, SellSerializer, SizeEntryDetailSerializer, SizeEntrySerializer,
    UserSerializer,
)

# The list's sort options, applied to whole brands instead of single size lines.
BRAND_ORDERING = {
    '-batch__date_added': ['-last_added'],
    'batch__date_added': ['last_added'],
    '-quantity': ['-pairs'],
    'quantity': ['pairs'],
    '-batch__bought_price': ['-max_price'],
    'batch__bought_price': ['min_price'],
}


def _size_key(size):
    """Sort sizes as numbers (9.5 before 40), with anything non-numeric after."""
    try:
        return (0, float(size), size)
    except ValueError:
        return (1, 0, size)


class LoginView(TokenObtainPairView):
    throttle_classes = [ScopedRateThrottle]
    throttle_scope = 'login'


class MeView(APIView):
    @extend_schema(responses=UserSerializer)
    def get(self, request):
        return Response(UserSerializer(request.user).data)


class BatchViewSet(mixins.CreateModelMixin, mixins.RetrieveModelMixin, mixins.UpdateModelMixin,
                   mixins.DestroyModelMixin, mixins.ListModelMixin, viewsets.GenericViewSet):
    """Deliveries. Creating one applies the restock rule (same brand + size + price tops up the existing line)."""

    queryset = Batch.objects.all()
    serializer_class = BatchSerializer
    parser_classes = [JSONParser, MultiPartParser, FormParser]
    ordering_fields = ['date_added', 'brand', 'bought_price']
    http_method_names = ['get', 'post', 'patch', 'delete']

    @extend_schema(request=BatchCreateSerializer, responses={201: RestockResultSerializer})
    def create(self, request):
        data = BatchCreateSerializer(data=request.data)
        data.is_valid(raise_exception=True)
        result = services.add_stock(
            data.validated_data['brand'],
            data.validated_data['bought_price'],
            [(row['size'], row['quantity']) for row in data.validated_data['sizes']],
            data.validated_data.get('picture'),
        )
        body = RestockResultSerializer(
            {'batch': result.batch, 'lines': [vars(line) for line in result.lines]},
            context={'request': request},
        )
        return Response(body.data, status=status.HTTP_201_CREATED)

    def destroy(self, request, *args, **kwargs):
        try:
            return super().destroy(request, *args, **kwargs)
        except ProtectedError:
            return Response({'detail': _('This delivery already has sales, so it can’t be deleted.')},
                            status=status.HTTP_409_CONFLICT)

    @action(detail=False)
    def brands(self, request):
        """Distinct brand names, for autocomplete on the add form."""
        return Response(list(Batch.objects.order_by('brand').values_list('brand', flat=True).distinct()))


@extend_schema(parameters=[OpenApiParameter('code', OpenApiTypes.STR, OpenApiParameter.PATH)])
class SizeEntryViewSet(mixins.RetrieveModelMixin, mixins.UpdateModelMixin, mixins.DestroyModelMixin,
                       mixins.ListModelMixin, viewsets.GenericViewSet):
    """One size line of a delivery: the thing a barcode points to."""

    queryset = SizeEntry.objects.select_related('batch')
    lookup_field = 'code'
    filterset_class = SizeEntryFilter
    ordering_fields = ['batch__date_added', 'quantity', 'batch__bought_price', 'batch__brand', 'size']
    ordering = ['-batch__date_added', 'batch__brand', 'size']
    http_method_names = ['get', 'post', 'patch', 'delete']

    def get_serializer_class(self):
        return SizeEntryDetailSerializer if self.action == 'retrieve' else SizeEntrySerializer

    def get_object(self):
        # Codes are case-insensitive for people typing them in from a label.
        self.kwargs[self.lookup_field] = self.kwargs[self.lookup_field].strip().upper()
        return super().get_object()

    def perform_update(self, serializer):
        # Correcting the count (a lost or miscounted pair) must not look like a sale.
        entry = serializer.instance
        delta = serializer.validated_data.get('quantity', entry.quantity) - entry.quantity
        with transaction.atomic():
            serializer.save(initial_quantity=max(0, entry.initial_quantity + delta))

    def destroy(self, request, *args, **kwargs):
        try:
            return super().destroy(request, *args, **kwargs)
        except ProtectedError:
            return Response({'detail': _('This size already has sales, so it can’t be deleted.')},
                            status=status.HTTP_409_CONFLICT)

    @extend_schema(responses=BrandGroupSerializer(many=True))
    @action(detail=False)
    def grouped(self, request):
        """The stock list one brand per row: same filters as the list, paginated by brand so a brand is never split."""
        entries = self.filter_queryset(self.get_queryset()).order_by()
        groups = (
            entries.values(brand=F('batch__brand'))
            .annotate(
                pairs=Sum('quantity'), min_price=Min('batch__bought_price'), max_price=Max('batch__bought_price'),
                last_added=Max('batch__date_added'), deliveries=Count('batch', distinct=True),
            )
            .order_by(*BRAND_ORDERING.get(request.query_params.get('ordering'), ['-last_added']), 'brand')
        )
        page = self.paginate_queryset(groups)

        by_brand = defaultdict(list)
        for entry in entries.filter(batch__brand__in=[group['brand'] for group in page]):
            by_brand[entry.batch.brand].append(entry)
        for group in page:
            group['entries'] = sorted(by_brand[group['brand']], key=lambda e: (_size_key(e.size), -e.batch.date_added.timestamp()))
        return self.get_paginated_response(BrandGroupSerializer(page, many=True, context={'request': request}).data)

    @extend_schema(request=SellSerializer, responses={201: SellResultSerializer, 409: OpenApiResponse(description='Sold out')})
    @action(detail=True, methods=['post'])
    def sell(self, request, code=None):
        entry = self.get_object()
        data = SellSerializer(data=request.data)
        data.is_valid(raise_exception=True)
        try:
            sale = services.sell_one(entry, data.validated_data['sold_price'])
        except services.OutOfStock:
            return Response({'detail': _('This size is sold out.')}, status=status.HTTP_409_CONFLICT)
        body = SellResultSerializer({'sale': sale, 'entry': entry}, context={'request': request})
        return Response(body.data, status=status.HTTP_201_CREATED)

    @extend_schema(responses={(200, 'image/png'): OpenApiTypes.BINARY})
    @action(detail=True)
    def label(self, request, code=None):
        entry = self.get_object()
        response = HttpResponse(services.label_png(entry), content_type='image/png')
        response['Content-Disposition'] = f'attachment; filename="label-{entry.code}.png"'
        return response


class SaleViewSet(mixins.ListModelMixin, mixins.RetrieveModelMixin, viewsets.GenericViewSet):
    """Sales history (read-only). A mistaken sale is removed in the Django admin, which restocks the pair."""

    queryset = Sale.objects.select_related('size_entry__batch')
    serializer_class = SaleSerializer
    filterset_class = SaleFilter
    ordering_fields = ['sold_at', 'sold_price']
    ordering = ['-sold_at']


class LabelSheetView(APIView):
    """A4 PDF of 50×30 mm barcode labels. Marks the lines as printed."""

    @extend_schema(request=LabelSheetSerializer, responses={(200, 'application/pdf'): OpenApiTypes.BINARY})
    def post(self, request):
        data = LabelSheetSerializer(data=request.data)
        data.is_valid(raise_exception=True)
        items = data.validated_data['items']
        codes = [item['code'].strip().upper() for item in items]
        entries = SizeEntry.objects.select_related('batch').in_bulk(codes, field_name='code')
        missing = sorted(set(codes) - entries.keys())
        if missing:
            raise ValidationError({'items': _('Unknown codes: %(codes)s') % {'codes': ', '.join(missing)}})

        sheet = [entries[code] for code, item in zip(codes, items) for _n in range(item['copies'])]
        pdf = services.labels_pdf(sheet)
        SizeEntry.objects.filter(code__in=codes).update(label_printed=True)
        response = HttpResponse(pdf, content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="labels-{timezone.localdate():%Y-%m-%d}.pdf"'
        return response


def _date_range(request, default_days):
    query = DateRangeSerializer(data=request.query_params)
    query.is_valid(raise_exception=True)
    end = query.validated_data.get('end') or timezone.localdate()
    start = query.validated_data.get('start') or end - timedelta(days=default_days - 1)
    return start, end, query.validated_data


class StatsView(APIView):
    @extend_schema(parameters=[DateRangeSerializer], responses=OpenApiTypes.OBJECT)
    def get(self, request):
        start, end, query = _date_range(request, default_days=30)
        sales = Sale.objects.filter(sold_at__date__gte=start, sold_at__date__lte=end)
        value, pairs = stats.inventory_value()
        step, points = stats.sales_over_time(sales, start, end)
        slow_days = query['slow_days']
        return Response({
            'start': start,
            'end': end,
            'inventory': {'value': value, 'pairs': pairs},
            'sales': stats.sales_summary(sales),
            'best_brands': stats.best_sellers(sales, 'size_entry__batch__brand'),
            'best_sizes': stats.best_sellers(sales, 'size_entry__size'),
            'slow_moving': {
                'days': slow_days,
                'entries': SizeEntrySerializer(stats.slow_moving(slow_days)[:30], many=True, context={'request': request}).data,
            },
            'timeline': {
                'step': step,
                'points': [{k: p[k] for k in ('date', 'revenue', 'profit', 'units')} for p in points],
            },
        })


class ExportView(APIView):
    """Download inventory or sales history as CSV or Excel."""

    @extend_schema(
        parameters=[DateRangeSerializer],
        responses={(200, 'text/csv'): OpenApiTypes.BINARY, (200, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'): OpenApiTypes.BINARY},
    )
    def get(self, request, kind, fmt):
        if kind == 'inventory':
            headers = [_('Code'), _('Brand'), _('Size'), _('Pairs left'), _('Pairs bought'), _('Bought price'), _('Stock value'), _('Date added')]
            rows = [
                [e.code, e.batch.brand, e.size, e.quantity, e.initial_quantity, e.batch.bought_price,
                 e.quantity * e.batch.bought_price, timezone.localtime(e.batch.date_added).strftime('%Y-%m-%d %H:%M')]
                for e in SizeEntry.objects.select_related('batch').order_by('batch__brand', 'size')
            ]
            name = f'inventory-{timezone.localdate():%Y-%m-%d}'
        else:
            start, end, _query = _date_range(request, default_days=365)
            headers = [_('Date'), _('Code'), _('Brand'), _('Size'), _('Bought price'), _('Sold price'), _('Profit')]
            sales = Sale.objects.select_related('size_entry__batch').filter(sold_at__date__gte=start, sold_at__date__lte=end)
            rows = [
                [timezone.localtime(s.sold_at).strftime('%Y-%m-%d %H:%M'), s.size_entry.code, s.size_entry.batch.brand,
                 s.size_entry.size, s.size_entry.batch.bought_price, s.sold_price, s.profit]
                for s in sales.order_by('sold_at')
            ]
            name = f'sales-{start:%Y-%m-%d}-{end:%Y-%m-%d}'

        if fmt == 'xlsx':
            from openpyxl import Workbook
            from openpyxl.styles import Font

            book = Workbook()
            sheet = book.active
            sheet.title = kind
            sheet.append(headers)
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            for row in rows:
                sheet.append(row)
            sheet.freeze_panes = 'A2'
            out = io.BytesIO()
            book.save(out)
            response = HttpResponse(out.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        else:
            response = HttpResponse(content_type='text/csv; charset=utf-8')
            response.write('﻿')  # BOM so Excel shows Cyrillic brand names correctly
            csv.writer(response).writerows([headers, *rows])
        response['Content-Disposition'] = f'attachment; filename="{name}.{fmt}"'
        return response


def _report_day(request):
    query = DayQuerySerializer(data=request.query_params)
    query.is_valid(raise_exception=True)
    return query.validated_data.get('date') or timezone.localdate()


class DailyReportListView(APIView):
    """Per-day totals (sold and received) for a range, newest day first. Defaults to the last 30 days."""

    @extend_schema(parameters=[DateRangeSerializer], responses=OpenApiTypes.OBJECT)
    def get(self, request):
        start, end, _query = _date_range(request, default_days=30)
        if (end - start).days > 366:
            raise ValidationError({'start': _('Pick a range of one year or less.')})
        return Response({'start': start, 'end': end, 'days': reports.days(start, end)})


class DailyReportView(APIView):
    """Everything that happened on one day. Defaults to today (shop time)."""

    @extend_schema(parameters=[DayQuerySerializer], responses=OpenApiTypes.OBJECT)
    def get(self, request):
        report = reports.daily(_report_day(request))
        report['sales_list'] = SaleSerializer(report['sales_list'], many=True).data
        return Response(report)


class DailyReportExportView(APIView):
    """One day's report as an Excel workbook: sales and received stock on separate sheets."""

    @extend_schema(parameters=[DayQuerySerializer], responses={(200, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'): OpenApiTypes.BINARY})
    def get(self, request):
        from openpyxl import Workbook
        from openpyxl.styles import Font

        day = _report_day(request)
        report = reports.daily(day)
        bold = Font(bold=True)
        book = Workbook()

        sales = book.active
        sales.title = _('Sales')
        sales.append([_('Time'), _('Code'), _('Brand'), _('Size'), _('Bought price'), _('Sold price'), _('Profit')])
        for sale in report['sales_list']:
            entry = sale.size_entry
            sales.append([timezone.localtime(sale.sold_at).strftime('%H:%M'), entry.code, entry.batch.brand, entry.size,
                          entry.batch.bought_price, sale.sold_price, sale.profit])
        totals = report['sales']
        sales.append([_('Total'), '', '', totals['units'], totals['cost'], totals['revenue'], totals['profit']])

        received = book.create_sheet(_('Received'))
        received.append([_('Code'), _('Brand'), _('Size'), _('Pairs'), _('Bought price'), _('Value')])
        for row in report['received']:
            received.append([row['code'], row['brand'], row['size'], row['quantity'], row['bought_price'],
                             row['quantity'] * row['bought_price']])
        received.append([_('Total'), '', '', report['received_summary']['pairs'], '', report['received_summary']['value']])

        for sheet in (sales, received):
            for cell in (*sheet[1], *sheet[sheet.max_row]):
                cell.font = bold
            sheet.freeze_panes = 'A2'

        out = io.BytesIO()
        book.save(out)
        response = HttpResponse(out.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="daily-report-{day:%Y-%m-%d}.xlsx"'
        return response
